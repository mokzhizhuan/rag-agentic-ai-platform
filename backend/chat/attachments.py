"""Safe, extract-only chat file intake. No arbitrary code execution or archives."""
import csv
import io
import json
import re
from pathlib import PurePath

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_FILES = 3
MAX_TOTAL_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 13000
TEXT_SUFFIXES = {'.txt','.md','.py','.js','.jsx','.ts','.tsx','.json','.jsonl','.html','.css','.sql','.yaml','.yml','.xml','.log','.toml','.ini','.cfg','.sh','.ps1','.java','.go','.rs','.c','.cpp','.h','.cs','.ipynb'}

class AttachmentError(ValueError):
    pass

def _name(name):
    safe = str(name or '').replace('\\','/').split('/')[-1].strip()
    if not safe or len(safe)>150 or safe.startswith('.'):
        raise AttachmentError('Invalid attachment filename')
    return safe

def _extract_ifc(filename, payload):
    # IFC-SPF (ISO 10303-21) text format: concise entity census, not 3D geometry.
    if not payload[:4096].lstrip().upper().startswith(b'ISO-10303-21;'):
        raise AttachmentError(f'{filename}: only text-based IFC-SPF is supported; IFCZIP and IFCXML require different parsers')
    decoded = payload.decode('latin-1')
    if 'DATA;' not in decoded.upper():
        raise AttachmentError(f'{filename}: no IFC DATA section found')
    import collections
    counts = collections.Counter(re.findall(r'(?im)^\s*#\d+\s*=\s*(IFC[A-Z0-9_]+)\s*\(', decoded))
    if not counts:
        raise AttachmentError(f'{filename}: no readable IFC entities found')
    # Include project, storeys, spaces, elements, materials and property relationships.
    priority = ('IFCPROJECT','IFCSITE','IFCBUILDING','IFCBUILDINGSTOREY','IFCSPACE',
                'IFCWALL','IFCWALLSTANDARDCASE','IFCDOOR','IFCWINDOW','IFCSLAB',
                'IFCBEAM','IFCCOLUMN','IFCSTAIR','IFCROOF','IFCMATERIAL',
                'IFCPROPERTYSET','IFCRELDEFINESBYPROPERTIES')
    names = [k for k in priority if k in counts]
    names += [k for k,_ in counts.most_common(45) if k not in names]
    lines = [f'IFC STEP building model: {filename}',
             f'Entity records: {sum(counts.values())}; distinct entity types: {len(counts)}.',
             'Entity type counts:']
    lines.extend(f'{name}: {counts[name]}' for name in names[:60])
    entities = re.findall(r'(?im)^\s*(#\d+\s*=\s*(?:IFCPROJECT|IFCSITE|IFCBUILDING|IFCBUILDINGSTOREY|IFCSPACE|IFCWALL|IFCDOOR|IFCWINDOW|IFCPROPERTYSET)\s*\([^;]{0,650}\);)', decoded)
    if entities:
        lines.append('Sample entity records (may contain identifiers and labels):')
        lines.extend(x.strip() for x in entities[:35])
    lines.append('LIMITATION: This is a lightweight IFC entity summary, not full BIM geometry, quantities, topology or validated property extraction. For advanced IFC analysis integrate IfcOpenShell.')
    return '\n'.join(lines)


def _extract_bmp(filename, payload):
    import struct
    if len(payload) < 54 or payload[:2] != b'BM':
        raise AttachmentError(f'{filename}: invalid BMP header')
    header_size = struct.unpack_from('<I', payload, 14)[0]
    if header_size < 40:
        raise AttachmentError(f'{filename}: unsupported BMP header variant')
    width, height = struct.unpack_from('<ii', payload, 18)
    planes, bit_depth = struct.unpack_from('<HH', payload, 26)
    compression = struct.unpack_from('<I', payload, 30)[0]
    if width <= 0 or height == 0 or abs(width) > 100000 or abs(height) > 100000 or planes != 1:
        raise AttachmentError(f'{filename}: invalid BMP dimensions')
    return (f'BMP image attachment: {filename}\nWidth: {width} pixels\nHeight: {abs(height)} pixels\n'
            f'Bits per pixel: {bit_depth}\nCompression code: {compression}\n'
            'LIMITATION: Only image metadata has been extracted. Image pixels and drawing contents have NOT been examined. '
            'Do not claim to see shapes, text, objects, dimensions or annotations in the bitmap. '
            'Use a vision-capable model and an image input pipeline for visual interpretation.')


def _extract(filename, payload):
    suffix = PurePath(filename).suffix.lower()
    if suffix == '.ifc':
        return _extract_ifc(filename, payload)
    if suffix == '.bmp':
        return _extract_bmp(filename, payload)
    if suffix in TEXT_SUFFIXES:
        try:
            decoded = payload.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise AttachmentError(f'{filename}: text must be UTF-8')
        if suffix == '.ipynb':
            try:
                notebook = json.loads(decoded)
                decoded = '\n'.join(''.join(c.get('source',[])) for c in notebook.get('cells',[]))
            except (ValueError,TypeError,AttributeError):
                raise AttachmentError(f'{filename}: invalid notebook')
        return decoded
    if suffix in ('.csv','.tsv'):
        try: raw=payload.decode('utf-8-sig')
        except UnicodeDecodeError: raise AttachmentError(f'{filename}: CSV/TSV must be UTF-8')
        reader=csv.reader(io.StringIO(raw), delimiter='\t' if suffix=='.tsv' else ',')
        return '\n'.join(' | '.join(row[:30])[:1200] for i,row in enumerate(reader) if i<180)
    if suffix in ('.xlsx','.xlsm'):
        from openpyxl import load_workbook
        try:
            wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=True)
            lines=[]
            for ws in wb.worksheets[:8]:
                lines.append('Sheet: '+ws.title)
                for i,row in enumerate(ws.iter_rows(values_only=True)):
                    if i>=120:break
                    lines.append(' | '.join(str(v)[:200] if v is not None else '' for v in row[:30]))
            wb.close()
            return '\n'.join(lines)
        except Exception as exc:
            raise AttachmentError(f'{filename}: invalid spreadsheet') from exc
    if suffix=='.docx':
        from docx import Document
        try:
            doc=Document(io.BytesIO(payload))
            return '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        except Exception as exc: raise AttachmentError(f'{filename}: invalid DOCX') from exc
    if suffix=='.pdf':
        from pypdf import PdfReader
        try:
            pdf=PdfReader(io.BytesIO(payload),strict=False)
            if len(pdf.pages)>80: raise AttachmentError('PDF exceeds 80-page limit')
            return '\n'.join((page.extract_text() or '') for page in pdf.pages)
        except AttachmentError:raise
        except Exception as exc:raise AttachmentError(f'{filename}: PDF extraction failed') from exc
    if suffix=='.pptx':
        from pptx import Presentation
        try:
            slides=Presentation(io.BytesIO(payload)).slides
            return '\n'.join(' | '.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in slides)
        except Exception as exc:raise AttachmentError(f'{filename}: invalid PPTX') from exc
    raise AttachmentError(f'{filename}: unsupported format {suffix or "(no extension)"}. Other images, video, audio and archives need separate processors.')

def extract_chat_files(files):
    if len(files)>MAX_FILES:raise AttachmentError('Maximum 3 files per message')
    total=0
    extracted=[]
    for f in files:
        name=_name(f.name)
        if f.size>MAX_FILE_BYTES:raise AttachmentError(f'{name}: maximum size is 5 MB')
        total+=f.size
        if total>MAX_TOTAL_BYTES:raise AttachmentError('Maximum combined size is 10 MB')
        data=f.read(MAX_FILE_BYTES+1)
        if len(data)>MAX_FILE_BYTES:raise AttachmentError(f'{name}: maximum size is 5 MB')
        text=_extract(name,data)
        if not text.strip():raise AttachmentError(f'{name}: no readable text was found (scanned files need OCR)')
        extracted.append({'filename':name,'text':text[:MAX_EXTRACTED_CHARS]})
    return extracted
