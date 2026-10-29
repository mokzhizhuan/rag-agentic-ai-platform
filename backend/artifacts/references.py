"""Bounded, in-memory text extraction for optional artifact reference documents.

No file is persisted, executed, or indexed. Image/video attachments are intentionally
unsupported until a multimodal reference adapter is connected.
"""
from io import BytesIO
from pathlib import Path

MAX_FILES = 3
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_REFERENCE_CHARS = 14000
TEXT_SUFFIXES = {'.txt', '.md', '.csv', '.json', '.yaml', '.yml', '.py', '.js', '.ts', '.tsx', '.html', '.css', '.sql'}


class ReferenceError(Exception):
    pass


def _extract(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix in TEXT_SUFFIXES:
        try:
            return data.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ReferenceError(f'{name}: text must be UTF-8 encoded.') from exc
    if suffix == '.pdf':
        from pypdf import PdfReader
        return '\n'.join((page.extract_text() or '') for page in PdfReader(BytesIO(data)).pages[:35])
    if suffix == '.docx':
        from docx import Document
        doc = Document(BytesIO(data))
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables[:12]:
            chunks.extend(' | '.join(cell.text for cell in row.cells) for row in table.rows[:80])
        return '\n'.join(chunks)
    if suffix == '.pptx':
        from pptx import Presentation
        prs = Presentation(BytesIO(data))
        return '\n'.join(shape.text for slide in prs.slides[:35] for shape in slide.shapes if shape.has_text_frame)
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        chunks = []
        try:
            for sheet in wb.worksheets[:8]:
                chunks.append(f'Sheet: {sheet.title}')
                for row in sheet.iter_rows(max_row=150, max_col=25, values_only=True):
                    chunks.append(' | '.join('' if v is None else str(v)[:250] for v in row))
        finally:
            wb.close()
        return '\n'.join(chunks)
    raise ReferenceError(f'{name}: unsupported reference type. Use text, PDF, DOCX, PPTX or XLSX files.')


def read_references(files) -> str:
    if not files:
        return ''
    if len(files) > MAX_FILES:
        raise ReferenceError('Maximum 3 reference files per generation request.')
    fragments = []
    available = MAX_REFERENCE_CHARS
    for file in files:
        if file.size > MAX_FILE_BYTES:
            raise ReferenceError(f'{file.name}: maximum file size is 5 MB.')
        if file.size == 0:
            raise ReferenceError(f'{file.name}: empty file.')
        try:
            text = _extract(file.name, file.read())
        except ReferenceError:
            raise
        except Exception as exc:
            raise ReferenceError(f'Could not extract {file.name}. Check that it is a valid file.') from exc
        if not text.strip():
            raise ReferenceError(f'{file.name}: no extractable text found. Scanned documents are not supported.')
        segment = f'\n--- REFERENCE: {Path(file.name).name[:100]} ---\n{text[:available]}\n'
        fragments.append(segment)
        available -= min(available, len(text))
        if available <= 0:
            break
    return ''.join(fragments)
