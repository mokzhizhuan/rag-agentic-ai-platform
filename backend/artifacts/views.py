import logging
import os
import re
import uuid
from django.core.files.base import ContentFile
from django.http import FileResponse
from django.shortcuts import get_object_or_404
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from chat.models import Conversation
from .generator import ArtifactGenerationError, create_artifact, FORMAT_EXT
from .models import GeneratedArtifact
from .references import ReferenceError
from .multimodal import prepare_references
from .generator import choose_format
from .video_builder import build_video, VideoBuildError
from .models import CharacterProfile
from .characters import ROLE_PROMPTS, LOCK_PROMPT

logger = logging.getLogger(__name__)

def serialize(a):
    return {'id': a.pk, 'filename': a.filename, 'format': a.format, 'byte_size': a.byte_size,
            'conversation_id': a.conversation_id, 'created_at': a.created_at.isoformat(),
            'download_path': f'/api/artifacts/{a.pk}/download/'}

@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def artifacts(request):
    owner_sub = request.user.sub
    if request.method == 'GET':
        items = GeneratedArtifact.objects.filter(owner_sub=owner_sub)
        conversation_id = request.query_params.get('conversation_id')
        if conversation_id:
            if not conversation_id.isdecimal():
                return Response({'detail': 'Invalid conversation_id'}, status=400)
            items = items.filter(conversation_id=int(conversation_id))
        return Response({'artifacts': [serialize(a) for a in items[:100]]})

    prompt = request.data.get('prompt', '')
    provider = request.data.get('provider', 'ollama')
    output_format = request.data.get('format', 'auto')
    cid = request.data.get('conversation_id')
    reference_role = str(request.data.get('reference_role','style'))
    character_lock = str(request.data.get('character_lock','false')).lower() in ('true','1','yes')
    character_id = str(request.data.get('character_id','')).strip()
    if reference_role not in ROLE_PROMPTS:
        return Response({'detail':'Unknown reference role.'},status=400)
    profile = None
    if character_id:
        if not character_id.isdecimal():
            return Response({'detail':'Invalid character profile ID.'},status=400)
        profile = get_object_or_404(CharacterProfile,pk=int(character_id),owner_sub=owner_sub)
    if character_lock and not (profile or any(f.name.lower().endswith(('.png','.jpg','.jpeg','.webp')) for f in request.FILES.getlist('files'))):
        return Response({'detail':'Character lock requires an image upload or saved character profile.'},status=400)
    identity_prompt = ROLE_PROMPTS['identity' if character_lock else reference_role]
    anchor_prompt = ('\nReference role: '+identity_prompt + ('\n'+LOCK_PROMPT if character_lock else '')) if (profile or request.FILES.getlist('files')) else ''
    if profile and profile.description:
        anchor_prompt += '\nCharacter profile '+profile.name+': '+profile.description
    if not isinstance(prompt, str) or not 5 <= len(prompt.strip()) <= 2500:
        return Response({'detail': 'Prompt must be 5–2500 characters.'}, status=400)
    if provider not in ('ollama', 'openai'):
        return Response({'detail': 'Provider must be ollama or openai.'}, status=400)
    if not isinstance(output_format, str) or output_format not in set(FORMAT_EXT) | {'auto', 'video', 'video_clip', 'video_reel'}:
        return Response({'detail': 'Unknown artifact format.'}, status=400)
    if provider == 'openai' and not os.getenv('OPENAI_API_KEY') and not (output_format in {'video','video_clip','video_reel'} and (request.FILES.getlist('files') or profile)):
        return Response({'detail': 'OpenAI API key is not configured.'}, status=400)
    conversation = None
    if cid not in (None, ''):
        if isinstance(cid, bool) or not str(cid).isdecimal():
            return Response({'detail': 'Invalid conversation_id.'}, status=400)
        conversation = get_object_or_404(Conversation, pk=int(cid), owner_sub=owner_sub)
    try:
        if output_format in {'video', 'video_clip', 'video_reel'}:
            video_mode = 'video_reel' if output_format == 'video' else output_format
            try:
                duration = int(request.data.get('duration', '30' if video_mode=='video_reel' else '8'))
            except (ValueError, TypeError):
                return Response({'detail': 'Invalid video duration.'}, status=400)
            files = request.FILES.getlist('files')
            if profile and files:
                return Response({'detail':'For motion video, choose either a saved character profile or uploaded still images (not both).'},status=400)
            if len(files)>3:
                return Response({'detail': 'Maximum 3 image references for videos.'}, status=400)
            stills = []
            for uploaded in files:
                if not uploaded.name.lower().endswith(('.png','.jpg','.jpeg','.webp')):
                    return Response({'detail': 'Video builder currently accepts PNG/JPG/WebP image references only.'}, status=400)
                stills.append(uploaded.read(8*1024*1024+1))
            if profile:
                with profile.image.open('rb') as handle:
                    stills.insert(0,handle.read(8*1024*1024+1))
            if not stills:
                if provider != 'openai':
                    return Response({'detail': 'Without reference images, choose OpenAI to generate a still image first, or attach an image with Ollama.'}, status=400)
                _, generated_still = create_artifact(prompt.strip()+anchor_prompt, 'openai', 'image')
                stills = [generated_still]
            fmt, binary = 'video', build_video(stills, video_mode, duration)
        else:
            effective_format = choose_format(prompt.strip(), output_format)
            uploads = request.FILES.getlist('files')
            if profile and len(uploads)>2:
                return Response({'detail':'Saved character anchor plus a maximum of two other references are allowed.'},status=400)
            generation_prompt, image_references = prepare_references(
                uploads, provider, effective_format, prompt.strip()+anchor_prompt)
            if profile:
                with profile.image.open('rb') as handle:
                    anchor_bytes=handle.read(8*1024*1024+1)
                if effective_format == 'image':
                    if provider != 'openai':
                        return Response({'detail':'Reference-guided image editing currently requires OpenAI.'},status=400)
                    image_references=[(profile.name,anchor_bytes)] + image_references[:2]
                else:
                    generation_prompt += '\nSaved character profile: '+profile.name+' ('+profile.description+'). Image anchor cannot be interpreted by a text-only generator.'
            fmt, binary = create_artifact(generation_prompt, provider, effective_format, image_references=image_references)
    except VideoBuildError as exc:
        return Response({'detail': str(exc)}, status=422)
    except ReferenceError as exc:
        return Response({'detail': str(exc)}, status=400)
    except ArtifactGenerationError as exc:
        return Response({'detail': str(exc)}, status=422)
    except Exception:
        logger.exception('Artifact generation failed')
        return Response({'detail': 'Generation failed. See backend logs for details.'}, status=503)
    if len(binary) > 15 * 1024 * 1024:
        return Response({'detail': 'Generated file exceeds the 15 MiB limit.'}, status=422)
    filename = f'generated-{uuid.uuid4().hex[:10]}.{"mp4" if fmt=="video" else FORMAT_EXT[fmt]}'
    item = GeneratedArtifact(owner_sub=owner_sub, conversation=conversation, prompt=prompt.strip(),
                             filename=filename, format=fmt, provider=provider, byte_size=len(binary))
    item.content.save(filename, ContentFile(binary), save=True)
    return Response({'artifact': serialize(item)}, status=201)

@api_view(['GET', 'DELETE'])
@permission_classes([IsAuthenticated])
def artifact_detail(request, pk):
    a = get_object_or_404(GeneratedArtifact, pk=pk, owner_sub=request.user.sub)
    if request.method == 'DELETE':
        a.content.delete(save=False)
        a.delete()
        return Response(status=204)
    return Response({'artifact': serialize(a)})

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def download_artifact(request, pk):
    a = get_object_or_404(GeneratedArtifact, pk=pk, owner_sub=request.user.sub)
    if not a.content:
        return Response({'detail':'Artifact file is missing.'}, status=404)
    try:
        return FileResponse(a.content.open('rb'), as_attachment=True, filename=a.filename,
                            content_type='application/octet-stream')
    except OSError:
        return Response({'detail':'Artifact file is unavailable.'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def preview_artifact(request, pk):
    """Preview text-only excerpts of Office files, authenticated and owner-scoped.

    Image/PDF/code previews use authenticated downloads on the frontend.
    This endpoint never returns executable HTML or raw office archives.
    """
    a = get_object_or_404(GeneratedArtifact, pk=pk, owner_sub=request.user.sub)
    if a.format not in {'word', 'excel', 'powerpoint'}:
        return Response({'detail': 'Use download endpoint for this file format.'}, status=400)
    if not a.content:
        return Response({'detail': 'Artifact file missing.'}, status=404)
    try:
        import io
        file_bytes = a.content.open('rb').read(15 * 1024 * 1024 + 1)
        if len(file_bytes) > 15 * 1024 * 1024:
            return Response({'detail': 'Preview size limit exceeded.'}, status=413)
        if a.format == 'word':
            from docx import Document
            d = Document(io.BytesIO(file_bytes))
            chunks = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables[:15]:
                for row in table.rows[:60]:
                    chunks.append(' | '.join(cell.text for cell in row.cells))
        elif a.format == 'excel':
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=False)
            chunks = []
            for ws in wb.worksheets[:8]:
                chunks.append('SHEET: ' + ws.title)
                for row in ws.iter_rows(max_row=100, max_col=25, values_only=True):
                    values = [str(v) if v is not None else '' for v in row]
                    if any(values): chunks.append(' | '.join(values))
            wb.close()
        else:
            from pptx import Presentation
            p = Presentation(io.BytesIO(file_bytes))
            chunks = []
            for n, slide in enumerate(p.slides[:40], 1):
                chunks.append('SLIDE ' + str(n))
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text.strip():
                        chunks.append(shape.text)
        preview = '\n'.join(chunks)
        return Response({'kind': 'text', 'text': preview[:60000], 'truncated': len(preview)>60000,
                         'note': 'Content outline only; layout, images, charts, and styles may not appear.'})
    except Exception:
        logger.exception('Artifact preview failed')
        return Response({'detail': 'This artifact could not be previewed. Download it to open in its native application.'}, status=422)
