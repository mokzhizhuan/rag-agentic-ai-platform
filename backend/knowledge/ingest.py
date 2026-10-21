from io import BytesIO
from pathlib import Path
from docx import Document as WordDocument
from pypdf import PdfReader
from .embeddings import embed, config
from .models import KnowledgeChunk


def extract(filename, binary):
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        reader = PdfReader(BytesIO(binary))
        return [(i+1, page.extract_text() or '') for i, page in enumerate(reader.pages)]
    if suffix == '.docx':
        doc = WordDocument(BytesIO(binary))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(' | '.join(c.text for c in row.cells))
        return [(None, '\n'.join(parts))]
    if suffix in ('.txt','.md'):
        return [(None,binary.decode('utf-8-sig'))]
    raise ValueError('Unsupported document format')


def split_text(text, size=1000, overlap=160):
    text=' '.join(text.split())
    if not text: return []
    chunks=[]; start=0
    while start < len(text):
        end=min(start+size,len(text))
        if end < len(text):
            boundary=text.rfind(' ',start+size//2,end)
            if boundary>start: end=boundary
        chunks.append(text[start:end])
        if end==len(text):break
        start=max(start+1,end-overlap)
    return chunks


def ingest_document(doc, binary):
    parts=[]
    for page,text in extract(doc.filename,binary):
        for chunk in split_text(text): parts.append((page,chunk))
    if not parts: raise ValueError('Document contains no extractable text (scanned PDFs require OCR)')
    if len(parts)>500: raise ValueError('Document exceeds 500 chunks in this starter version')
    provider, model=config()
    idx=0
    for batch_start in range(0,len(parts),12):
        batch=parts[batch_start:batch_start+12]
        vectors=embed([text for _,text in batch])
        KnowledgeChunk.objects.bulk_create([
            KnowledgeChunk(document=doc,index=batch_start+i,page=page,text=text,embedding=vector)
            for i,((page,text),vector) in enumerate(zip(batch,vectors))
        ])
        idx+=len(batch)
    doc.chunk_count=idx;doc.embedding_provider=provider;doc.embedding_model=model
    doc.status='ready';doc.error=''
    doc.save(update_fields=['chunk_count','embedding_provider','embedding_model','status','error'])
