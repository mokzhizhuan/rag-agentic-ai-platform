"""Safe document renderers and structural QA for generated report sections.
No generated content is executed; model text is treated as untrusted data.
"""
import io
import re
import zipfile
from xml.sax.saxutils import escape

class DocumentQualityError(ValueError):
    pass

_HEADING = re.compile(r'^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$')
_NUMBERED = re.compile(r'^\s*(?:\d+(?:\.\d+)*[.)]?\s+)(.+)$')
_BULLET = re.compile(r'^\s*[-*+]\s+(.+)$')
_TABLE_DIVIDER = re.compile(r'^\s*\|?[\s:\-|]+\|?\s*$')


def plain_heading(value):
    value = str(value or '').strip()
    match = _HEADING.match(value)
    if match:
        value = match.group(1).strip()
    value = re.sub(r'\*\*(.*?)\*\*', r'\1', value)
    return value.strip()


def normalize_sections(title, sections):
    title = plain_heading(title)
    if not title or not isinstance(sections, (tuple, list)) or not sections:
        raise DocumentQualityError('Report requires a title and nonempty sections.')
    result, known = [], set()
    for index, section in enumerate(sections, 1):
        if not isinstance(section, dict):
            raise DocumentQualityError(f'Section {index} must be an object.')
        heading = plain_heading(section.get('heading', ''))
        body = str(section.get('body') or '').strip()
        if not heading or len(body) < 60:
            raise DocumentQualityError(f'Section {index} is missing a heading or substantive body.')
        normalized_key = re.sub(r'^\d+(?:\.\d+)*[.)]?\s*', '', heading).casefold()
        if normalized_key in known:
            raise DocumentQualityError(f'Duplicate section heading: {heading}')
        known.add(normalized_key)
        # Remove an AI-duplicated section heading at the start of the body.
        lines = body.splitlines()
        while lines and not lines[0].strip():
            lines.pop(0)
        if lines:
            first = plain_heading(lines[0])
            first = re.sub(r'^\d+(?:\.\d+)*[.)]?\s*', '', first).casefold()
            if first == normalized_key:
                lines.pop(0)
                body = '\n'.join(lines).lstrip()
        if len(body) < 45:
            raise DocumentQualityError(f'Section {index} has insufficient body content.')
        result.append({'heading': heading, 'body': body})
    return title, result


def _blocks(body):
    """Parse conservative Markdown blocks. Does not interpret arbitrary HTML."""
    lines = body.splitlines()
    output, para, i = [], [], 0
    def flush():
        if para:
            output.append(('paragraph', ' '.join(para).strip()))
            para.clear()
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            flush(); i += 1; continue
        if line.startswith('```'):
            flush(); lang=line[3:].strip(); code=[]; i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i+=1
            output.append(('code', '\n'.join(code)[:14000], lang)); i += (i<len(lines)); continue
        heading = _HEADING.match(line)
        if heading:
            flush(); output.append(('subheading', plain_heading(heading.group(1)))); i+=1; continue
        if '|' in line and i+1<len(lines) and _TABLE_DIVIDER.match(lines[i+1]) and '|' in lines[i+1]:
            flush()
            table=[_cells(line)]; i+=2
            while i<len(lines) and '|' in lines[i] and lines[i].strip():
                table.append(_cells(lines[i])); i+=1
            width=len(table[0]); output.append(('table', [r[:width]+['']*max(0,width-len(r)) for r in table][:60])); continue
        match=_BULLET.match(line)
        if match:
            flush(); output.append(('bullet', match.group(1))); i+=1; continue
        match=_NUMBERED.match(line)
        if match and len(line)<250:
            flush(); output.append(('number', line)); i+=1; continue
        para.append(line); i+=1
    flush()
    return output


def _cells(line):
    return [part.strip() for part in line.strip().strip('|').split('|')]


def _inline_text(value):
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', str(value))
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text


def markdown_document(title, sections):
    title, sections = normalize_sections(title, sections)
    lines=['# '+title, '']
    for section in sections:
        lines += ['## '+section['heading'], '', section['body'].strip(), '']
    result='\n'.join(lines).rstrip()+'\n'
    if result.count('\n## ')!=len(sections):
        # A report section is allowed to contain third-level headings, but a
        # full second-level heading inside its body can be ambiguous.
        pass
    return result.encode('utf-8')


def render_word(title, sections):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    title, sections = normalize_sections(title, sections)
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.7)
    sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
    normal=doc.styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7)
    for name,size,color in [('Title',21,(24,55,90)),('Heading 1',15,(24,65,102)),('Heading 2',12,(34,78,111))]:
        style=doc.styles[name]; style.font.name='Aptos'; style.font.size=Pt(size)
        style.font.bold=True; style.font.color.rgb=RGBColor(*color)
    p=doc.add_paragraph(style='Title');p.add_run(title)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Technical Report').alignment=WD_ALIGN_PARAGRAPH.CENTER
    for section in sections:
        doc.add_heading(section['heading'],level=1)
        for block in _blocks(section['body']):
            kind=block[0]; value=block[1]
            if kind=='subheading': doc.add_heading(_inline_text(value),level=2)
            elif kind=='table':
                rows=value
                if rows and rows[0]:
                    tbl=doc.add_table(rows=1,cols=len(rows[0]));tbl.style='Light Shading Accent 1';tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
                    for j,c in enumerate(rows[0]):tbl.rows[0].cells[j].text=_inline_text(c)
                    for row in rows[1:]:
                        cells=tbl.add_row().cells
                        for j,c in enumerate(row):cells[j].text=_inline_text(c)
            elif kind=='code':
                para=doc.add_paragraph(style='Normal');run=para.add_run(value);run.font.name='Consolas';run.font.size=Pt(8)
            else:
                para=doc.add_paragraph(_inline_text(value),style='List Bullet' if kind=='bullet' else 'Normal')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Generated report • Review before operational use')
    buf=io.BytesIO();doc.save(buf); payload=buf.getvalue()
    if not zipfile.is_zipfile(io.BytesIO(payload)):
        raise DocumentQualityError('Invalid DOCX container.')
    check=Document(io.BytesIO(payload))
    headings=[p.text for p in check.paragraphs if p.style.name=='Heading 1']
    if headings != [s['heading'] for s in sections]:
        raise DocumentQualityError('DOCX heading verification failed.')
    return payload


def render_pdf(title, sections):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted, KeepTogether
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    title, sections=normalize_sections(title, sections)
    buf=io.BytesIO(); styles=getSampleStyleSheet()
    title_style=ParagraphStyle('DQTitle',parent=styles['Title'],fontSize=19,leading=24,textColor=colors.HexColor('#193D60'),spaceAfter=16,alignment=TA_CENTER)
    heading_style=ParagraphStyle('DQHeading',parent=styles['Heading2'],fontSize=13,leading=17,spaceBefore=15,spaceAfter=8,textColor=colors.HexColor('#215778'),keepWithNext=True)
    sub_style=ParagraphStyle('DQSub',parent=styles['Heading3'],fontSize=11,leading=15,spaceBefore=9,spaceAfter=5,keepWithNext=True)
    body_style=ParagraphStyle('DQBody',parent=styles['BodyText'],fontSize=9.2,leading=13.8,spaceAfter=8)
    bullet_style=ParagraphStyle('DQBullet',parent=body_style,leftIndent=14,firstLineIndent=-9)
    cell_style=ParagraphStyle('DQCell',parent=body_style,fontSize=7.3,leading=10,spaceAfter=0)
    story=[Paragraph(escape(title),title_style), Spacer(1,6)]
    for section in sections:
        story.append(Paragraph(escape(section['heading']),heading_style))
        for block in _blocks(section['body']):
            kind,value=block[0],block[1]
            if kind=='subheading': story.append(Paragraph(escape(_inline_text(value)),sub_style))
            elif kind=='table':
                rows=value
                if rows and rows[0]:
                    data=[[Paragraph(escape(_inline_text(str(v)[:800])),cell_style) for v in row] for row in rows]
                    table=Table(data,repeatRows=1,hAlign='LEFT',splitByRow=1)
                    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DFEAF3')),('LINEBELOW',(0,0),(-1,0),0.6,colors.HexColor('#7292AC')),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F5F8FC')])]))
                    story.append(table);story.append(Spacer(1,9))
            elif kind=='code':
                story.append(Preformatted(value[:7000],cell_style,maxLineLength=85,splitChars=' ,;'))
            elif kind=='bullet':story.append(Paragraph('&#8226; '+escape(_inline_text(value)),bullet_style))
            else:story.append(Paragraph(escape(_inline_text(value)).replace('\n','<br/>'),body_style))
    def decorate(canvas, document):
        canvas.saveState();canvas.setStrokeColor(colors.HexColor('#CBD9E5'));canvas.line(20*mm,19*mm,190*mm,19*mm)
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#586B7C'))
        canvas.drawString(20*mm,14*mm,'Generated technical report')
        canvas.drawRightString(190*mm,14*mm,f'Page {document.page}')
        canvas.restoreState()
    SimpleDocTemplate(buf,pagesize=A4,leftMargin=20*mm,rightMargin=20*mm,topMargin=20*mm,bottomMargin=24*mm,title=title).build(story,onFirstPage=decorate,onLaterPages=decorate)
    payload=buf.getvalue()
    if not payload.startswith(b'%PDF-') or b'%%EOF' not in payload[-2048:]:
        raise DocumentQualityError('Generated PDF failed structural checks.')
    try:
        from pypdf import PdfReader
    except ImportError:
        pass
    else:
        reader=PdfReader(io.BytesIO(payload))
        if len(reader.pages)<1:
            raise DocumentQualityError('Generated PDF contains no pages.')
    return payload
